from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "output" / "attacklab-zh" / "attacklab-zh.md"
OUTPUT_DOCX = ROOT / "output" / "attacklab-zh" / "attacklab-zh.docx"
ASSET_DIR = ROOT / "output" / "attacklab-zh" / "assets"

FONT_CJK = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_CJK_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")
FONT_MONO = Path(r"C:\Windows\Fonts\consola.ttf")
FONT_MONO_BOLD = Path(r"C:\Windows\Fonts\consolab.ttf")

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
TEXT = RGBColor(0x20, 0x20, 0x20)
MUTED = RGBColor(0x66, 0x66, 0x66)


def pil_font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size=size)


def centered_text(draw: ImageDraw.ImageDraw, box, text, font, fill="#111111"):
    left, top, right, bottom = box
    bounds = draw.textbbox((0, 0), text, font=font)
    width = bounds[2] - bounds[0]
    height = bounds[3] - bounds[1]
    draw.text(
        ((left + right - width) / 2, (top + bottom - height) / 2 - bounds[1]),
        text,
        font=font,
        fill=fill,
    )


def draw_arrow(draw, start, end, width=6, fill="#111111"):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    angle = __import__("math").atan2(y2 - y1, x2 - x1)
    length = 24
    spread = 0.55
    p1 = (
        x2 - length * __import__("math").cos(angle - spread),
        y2 - length * __import__("math").sin(angle - spread),
    )
    p2 = (
        x2 - length * __import__("math").cos(angle + spread),
        y2 - length * __import__("math").sin(angle + spread),
    )
    draw.polygon([end, p1, p2], fill=fill)


def create_figure_2(path: Path):
    image = Image.new("RGB", (1900, 820), "white")
    draw = ImageDraw.Draw(image)
    cjk = pil_font(FONT_CJK, 44)
    code = pil_font(FONT_MONO, 40)
    code_bold = pil_font(FONT_MONO_BOLD, 40)

    stack_left, stack_top, stack_right, stack_bottom = 430, 150, 720, 660
    draw.rectangle((stack_left, stack_top, stack_right, stack_bottom), outline="#111111", width=5)
    for y in (260, 440, 550):
        draw.line((stack_left, y, stack_right, y), fill="#111111", width=4)
    centered_text(draw, (stack_left, 70, stack_right, 145), "栈", cjk)
    centered_text(draw, (stack_left, 275, stack_right, 410), "·\n·\n·", cjk)

    points = [(575, 205), (575, 495), (575, 605)]
    for point in points:
        draw.ellipse((point[0] - 13, point[1] - 13, point[0] + 13, point[1] + 13), fill="#111111")

    gadget_boxes = [
        (980, 120, 1540, 240, "Gadget n 代码"),
        (980, 360, 1540, 480, "Gadget 2 代码"),
        (980, 570, 1540, 690, "Gadget 1 代码"),
    ]
    for left, top, right, bottom, label in gadget_boxes:
        draw.rectangle((left, top, right, bottom), outline="#111111", width=5)
        split = right - 105
        draw.line((split, top, split, bottom), fill="#111111", width=4)
        centered_text(draw, (left, top, split, bottom), label, cjk)
        centered_text(draw, (split, top, right, bottom), "c3", code)

    draw_arrow(draw, points[0], (980, 180))
    draw_arrow(draw, points[1], (980, 420))
    draw_arrow(draw, points[2], (980, 630))
    draw_arrow(draw, (260, 605), (430, 605))
    centered_text(draw, (95, 560, 260, 650), "%rsp", code_bold)
    image.save(path, dpi=(300, 300))


def draw_grid_table(draw, x, y, widths, row_height, rows, fonts, fills=None):
    total_width = sum(widths)
    for row_index, row in enumerate(rows):
        top = y + row_index * row_height
        left = x
        for col_index, cell in enumerate(row):
            right = left + widths[col_index]
            if fills and row_index < len(fills) and fills[row_index]:
                draw.rectangle((left, top, right, top + row_height), fill=fills[row_index])
            draw.rectangle((left, top, right, top + row_height), outline="#222222", width=3)
            centered_text(
                draw,
                (left + 4, top + 2, right - 4, top + row_height - 2),
                cell,
                fonts[row_index] if isinstance(fonts, list) else fonts,
            )
            left = right
    return y + len(rows) * row_height, total_width


def create_figure_3(path: Path):
    width, height = 2520, 2200
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    cjk = pil_font(FONT_CJK, 36)
    cjk_bold = pil_font(FONT_CJK_BOLD, 38)
    mono = pil_font(FONT_MONO, 31)
    mono_bold = pil_font(FONT_MONO_BOLD, 32)
    x = 70
    y = 45

    movq_rows = [
        ["源 S / 目的 D", "%rax", "%rcx", "%rdx", "%rbx", "%rsp", "%rbp", "%rsi", "%rdi"],
        ["%rax", "48 89 c0", "48 89 c1", "48 89 c2", "48 89 c3", "48 89 c4", "48 89 c5", "48 89 c6", "48 89 c7"],
        ["%rcx", "48 89 c8", "48 89 c9", "48 89 ca", "48 89 cb", "48 89 cc", "48 89 cd", "48 89 ce", "48 89 cf"],
        ["%rdx", "48 89 d0", "48 89 d1", "48 89 d2", "48 89 d3", "48 89 d4", "48 89 d5", "48 89 d6", "48 89 d7"],
        ["%rbx", "48 89 d8", "48 89 d9", "48 89 da", "48 89 db", "48 89 dc", "48 89 dd", "48 89 de", "48 89 df"],
        ["%rsp", "48 89 e0", "48 89 e1", "48 89 e2", "48 89 e3", "48 89 e4", "48 89 e5", "48 89 e6", "48 89 e7"],
        ["%rbp", "48 89 e8", "48 89 e9", "48 89 ea", "48 89 eb", "48 89 ec", "48 89 ed", "48 89 ee", "48 89 ef"],
        ["%rsi", "48 89 f0", "48 89 f1", "48 89 f2", "48 89 f3", "48 89 f4", "48 89 f5", "48 89 f6", "48 89 f7"],
        ["%rdi", "48 89 f8", "48 89 f9", "48 89 fa", "48 89 fb", "48 89 fc", "48 89 fd", "48 89 fe", "48 89 ff"],
    ]
    draw.text((x, y), "A. movq 指令的编码", font=cjk_bold, fill="#111111")
    y += 58
    draw.text((x + 30, y), "movq S, D", font=mono_bold, fill="#111111")
    y += 54
    y, _ = draw_grid_table(draw, x, y, [260] + [270] * 8, 56, movq_rows, [cjk_bold] + [mono] * 8, ["#E8EEF5"])

    y += 42
    draw.text((x, y), "B. popq 指令的编码", font=cjk_bold, fill="#111111")
    y += 62
    pop_rows = [
        ["操作 / 寄存器 R", "%rax", "%rcx", "%rdx", "%rbx", "%rsp", "%rbp", "%rsi", "%rdi"],
        ["popq R", "58", "59", "5a", "5b", "5c", "5d", "5e", "5f"],
    ]
    y, _ = draw_grid_table(draw, 330, y, [350] + [185] * 8, 60, pop_rows, [cjk_bold, mono], ["#E8EEF5"])

    y += 42
    draw.text((x, y), "C. movl 指令的编码", font=cjk_bold, fill="#111111")
    y += 58
    draw.text((x + 30, y), "movl S, D", font=mono_bold, fill="#111111")
    y += 54
    movl_rows = [
        ["源 S / 目的 D", "%eax", "%ecx", "%edx", "%ebx", "%esp", "%ebp", "%esi", "%edi"],
        ["%eax", "89 c0", "89 c1", "89 c2", "89 c3", "89 c4", "89 c5", "89 c6", "89 c7"],
        ["%ecx", "89 c8", "89 c9", "89 ca", "89 cb", "89 cc", "89 cd", "89 ce", "89 cf"],
        ["%edx", "89 d0", "89 d1", "89 d2", "89 d3", "89 d4", "89 d5", "89 d6", "89 d7"],
        ["%ebx", "89 d8", "89 d9", "89 da", "89 db", "89 dc", "89 dd", "89 de", "89 df"],
        ["%esp", "89 e0", "89 e1", "89 e2", "89 e3", "89 e4", "89 e5", "89 e6", "89 e7"],
        ["%ebp", "89 e8", "89 e9", "89 ea", "89 eb", "89 ec", "89 ed", "89 ee", "89 ef"],
        ["%esi", "89 f0", "89 f1", "89 f2", "89 f3", "89 f4", "89 f5", "89 f6", "89 f7"],
        ["%edi", "89 f8", "89 f9", "89 fa", "89 fb", "89 fc", "89 fd", "89 fe", "89 ff"],
    ]
    y, _ = draw_grid_table(draw, x, y, [260] + [270] * 8, 56, movl_rows, [cjk_bold] + [mono] * 8, ["#E8EEF5"])

    y += 42
    draw.text((x, y), "D. 双字节功能性 nop 指令的编码", font=cjk_bold, fill="#111111")
    y += 65
    nop_rows = [
        ["操作 / 寄存器 R", "%al", "%cl", "%dl", "%bl"],
        ["andb R, R", "20 c0", "20 c9", "20 d2", "20 db"],
        ["orb R, R", "08 c0", "08 c9", "08 d2", "08 db"],
        ["cmpb R, R", "38 c0", "38 c9", "38 d2", "38 db"],
        ["testb R, R", "84 c0", "84 c9", "84 d2", "84 db"],
    ]
    draw_grid_table(draw, 520, y, [460] + [260] * 4, 62, nop_rows, [cjk_bold] + [mono] * 4, ["#E8EEF5"])
    image.save(path, dpi=(300, 300))


def set_run_font(run, ascii_font="Calibri", east_asia="Microsoft YaHei", size=None, bold=None, italic=None, color=None):
    run.font.name = ascii_font
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade(element, fill):
    p_pr = element.get_or_add_pPr() if element.tag == qn("w:p") else element.get_or_add_tcPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name, indent, hanging in (
        ("List Bullet", 0.375, 0.188),
        ("List Bullet 2", 0.75, 0.188),
        ("List Number", 0.375, 0.188),
        ("List Number 2", 0.75, 0.188),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(indent)
        style.paragraph_format.first_line_indent = Inches(-hanging)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(paragraph, text):
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, ascii_font="Consolas", east_asia="Microsoft YaHei", size=10.3, color=DARK_BLUE)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=11, bold=True, color=TEXT)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=10.5, italic=True, color=MUTED)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=11, color=TEXT)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    shade(p._p, "F6F7F9")
    for index, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, ascii_font="Consolas", east_asia="Microsoft YaHei", size=9.2, color=TEXT)
        if index != len(lines) - 1:
            run.add_break()


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "2E74B5")
    borders.append(left)
    p_pr.append(borders)
    shade(p._p, "F4F6F9")
    add_inline(p, text)


def add_markdown_table(doc, lines):
    rows = []
    for index, line in enumerate(lines):
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if index == 1 and all(re.fullmatch(r":?-+:?", cell) for cell in cells):
            continue
        rows.append(cells)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    widths = [1152, 1728, 1152, 1440, 2448, 1440]
    set_table_geometry(table, widths)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            add_inline(p, value)
            if row_index == 0:
                shade(cell._tc, "E8EEF5")
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_image(doc, relative_path):
    path = SOURCE_MD.parent / relative_path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.35))


def build_docx():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_figure_2(ASSET_DIR / "figure-2-rop-gadgets-zh.png")
    create_figure_3(ASSET_DIR / "figure-3-instruction-encodings-zh.png")

    doc = Document()
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()
    first_h1 = next(i for i, line in enumerate(lines) if line.startswith("# "))
    second_h1 = next(i for i in range(first_h1 + 1, len(lines)) if lines[i].startswith("# "))

    course = doc.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course.paragraph_format.space_before = Pt(64)
    course.paragraph_format.space_after = Pt(8)
    run = course.add_run(lines[first_h1][2:].strip())
    set_run_font(run, size=12, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    run = title.add_run(lines[second_h1][2:].strip())
    set_run_font(run, size=22, bold=True, color=TEXT)

    i = second_h1 + 1
    while i < len(lines) and not lines[i].startswith("## "):
        text = lines[i].rstrip().rstrip("  ")
        if text:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)
        i += 1
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(16)

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].strip("\n"))
                i += 1
            add_code_block(doc, code)
            i += 1
            continue
        if stripped.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", stripped)
            if match:
                add_image(doc, match.group(1))
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        if stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, stripped[3:])
            i += 1
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, stripped[4:])
            i += 1
            continue
        if stripped.startswith("> "):
            add_callout(doc, stripped[2:])
            i += 1
            continue
        bullet = re.match(r"^(\s*)-\s+(.*)", line)
        if bullet:
            level = 1 if len(bullet.group(1)) >= 2 else 0
            p = doc.add_paragraph(style="List Bullet 2" if level else "List Bullet")
            add_inline(p, bullet.group(2))
            i += 1
            continue
        number = re.match(r"^(\s*)\d+\.\s+(.*)", line)
        if number:
            level = 1 if len(number.group(1)) >= 2 else 0
            p = doc.add_paragraph(style="List Number 2" if level else "List Number")
            add_inline(p, number.group(2))
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            candidate = lines[i]
            candidate_stripped = candidate.strip()
            if not candidate_stripped:
                break
            if (
                candidate_stripped.startswith(("## ", "### ", "```", "![", "|", "> "))
                or re.match(r"^\s*-\s+", candidate)
                or re.match(r"^\s*\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate_stripped)
            i += 1
        p = doc.add_paragraph()
        text = " ".join(part.rstrip().rstrip("  ") for part in paragraph_lines)
        add_inline(p, text)

    doc.core_properties.title = "攻击实验：理解缓冲区溢出漏洞（中文完整翻译）"
    doc.core_properties.subject = "CS:APP Attack Lab 中文翻译"
    doc.core_properties.author = "OpenAI Codex"
    doc.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT_DOCX)
